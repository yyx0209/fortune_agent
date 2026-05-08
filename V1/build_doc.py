from typing import List, Dict, Any
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH


def _init_doc(title: str, subtitle: str | None = None) -> Document:
    doc = Document()

    section = doc.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.4)

    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10.5)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title)
    run.bold = True
    run.font.size = Pt(18)

    if subtitle:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run(subtitle).italic = True

    doc.add_paragraph()
    return doc


def build_bazi_debate_docx(
    output_path: str,
    source_prompt: str,
    transcript: List[Dict[str, Any]],
    consensus_history: List[Dict[str, Any]],
    running_consensus: List[str],
    unresolved_points: List[str],
) -> None:
    """
    只输出辩论过程文档
    """
    doc = _init_doc(
        title="八字多 Agent 辩论过程记录",
        subtitle="A. 多轮辩论过程",
    )

    meta = doc.add_table(rows=3, cols=2)
    meta.style = "Table Grid"
    meta.cell(0, 0).text = "分析模式"
    meta.cell(0, 1).text = "多 agent 辩论 + 共识整合"
    meta.cell(1, 0).text = "当前累计共识数"
    meta.cell(1, 1).text = str(len(running_consensus))
    meta.cell(2, 0).text = "剩余未解分歧数"
    meta.cell(2, 1).text = str(len(unresolved_points))

    doc.add_paragraph()
    doc.add_heading("原始输入", level=1)
    doc.add_paragraph(source_prompt)

    doc.add_page_break()

    doc.add_heading("多轮辩论过程", level=1)
    doc.add_paragraph("以下内容按轮次记录各 agent 的发言，以及每轮结束后的共识整合结果。")

    rounds = sorted(set(x["round"] for x in transcript))
    for rd in rounds:
        doc.add_heading(f"第 {rd} 轮", level=2)

        round_records = [x for x in transcript if x["round"] == rd]
        for rec in round_records:
            doc.add_heading(rec["speaker"], level=3)

            info = doc.add_table(rows=2, cols=2)
            info.style = "Table Grid"
            info.cell(0, 0).text = "模型"
            info.cell(0, 1).text = rec.get("model", "")
            info.cell(1, 0).text = "核心判断"
            info.cell(1, 1).text = "\n".join([f"- {x}" for x in rec.get("core_claims", [])]) or "无"

            if rec.get("critique_of_previous"):
                doc.add_paragraph("对上一位分析者的逻辑批评：")
                for x in rec["critique_of_previous"]:
                    doc.add_paragraph(x, style="List Bullet")

            if rec.get("evidence"):
                doc.add_paragraph("主要证据：")
                for x in rec["evidence"]:
                    doc.add_paragraph(x, style="List Bullet")

            if rec.get("tentative_consensus"):
                doc.add_paragraph("本发言提出的可共识结论：")
                for x in rec["tentative_consensus"]:
                    doc.add_paragraph(x, style="List Bullet")

            if rec.get("uncertainties"):
                doc.add_paragraph("本发言承认的不确定点：")
                for x in rec["uncertainties"]:
                    doc.add_paragraph(x, style="List Bullet")

            doc.add_paragraph("完整发言：")
            doc.add_paragraph(rec.get("full_response", ""))

        round_cons = next((x for x in consensus_history if x["round"] == rd), None)
        if round_cons:
            doc.add_heading("本轮共识整合", level=3)
            t = doc.add_table(rows=4, cols=2)
            t.style = "Table Grid"
            t.cell(0, 0).text = "本轮总结"
            t.cell(0, 1).text = round_cons.get("summary", "")
            t.cell(1, 0).text = "本轮新增/强化共识"
            t.cell(1, 1).text = "\n".join([f"- {x}" for x in round_cons.get("agreed_points", [])]) or "无"
            t.cell(2, 0).text = "仍未解决分歧"
            t.cell(2, 1).text = "\n".join([f"- {x}" for x in round_cons.get("still_unresolved", [])]) or "无"
            t.cell(3, 0).text = "下一轮焦点"
            t.cell(3, 1).text = "\n".join([f"- {x}" for x in round_cons.get("next_focus", [])]) or "无"

        doc.add_paragraph()

    doc.save(output_path)


def build_bazi_result_docx(
    output_path: str,
    source_prompt: str,
    final_scores: Dict[str, Any],
    final_report: str,
    running_consensus: List[str],
    unresolved_points: List[str],
) -> None:
    """
    只输出最终结果文档
    """
    doc = _init_doc(
        title="八字多 Agent 最终共识结果",
        subtitle="B. 最终结果",
    )

    meta = doc.add_table(rows=3, cols=2)
    meta.style = "Table Grid"
    meta.cell(0, 0).text = "分析模式"
    meta.cell(0, 1).text = "多 agent 辩论后结果汇总"
    meta.cell(1, 0).text = "最终共识数"
    meta.cell(1, 1).text = str(len(running_consensus))
    meta.cell(2, 0).text = "剩余未解分歧数"
    meta.cell(2, 1).text = str(len(unresolved_points))

    doc.add_paragraph()
    doc.add_heading("原始输入", level=1)
    doc.add_paragraph(source_prompt)

    doc.add_page_break()

    doc.add_heading("最终累计共识", level=1)
    if running_consensus:
        for x in running_consensus:
            doc.add_paragraph(x, style="List Bullet")
    else:
        doc.add_paragraph("无")

    doc.add_heading("仍保留的少量分歧", level=1)
    if unresolved_points:
        for x in unresolved_points:
            doc.add_paragraph(x, style="List Bullet")
    else:
        doc.add_paragraph("无明显保留分歧。")

    doc.add_heading("命主人群位置评分表（满分10分）", level=1)
    score_table = doc.add_table(rows=1, cols=3)
    score_table.style = "Table Grid"
    score_table.rows[0].cells[0].text = "维度"
    score_table.rows[0].cells[1].text = "分数"
    score_table.rows[0].cells[2].text = "依据"

    for dim, val in final_scores.items():
        row = score_table.add_row().cells
        row[0].text = dim
        row[1].text = str(val.get("score", ""))
        row[2].text = val.get("reason", "")

    doc.add_heading("面向用户的最终报告", level=1)
    for para in final_report.split("\n"):
        para = para.strip()
        if para:
            doc.add_paragraph(para)

    doc.save(output_path)


def build_bazi_docx_files(
    debate_output_path: str,
    result_output_path: str,
    source_prompt: str,
    transcript: List[Dict[str, Any]],
    consensus_history: List[Dict[str, Any]],
    final_scores: Dict[str, Any],
    final_report: str,
    running_consensus: List[str],
    unresolved_points: List[str],
) -> None:
    """
    同时生成两个文件：
    1. 辩论过程
    2. 最终结果
    """
    build_bazi_debate_docx(
        output_path=debate_output_path,
        source_prompt=source_prompt,
        transcript=transcript,
        consensus_history=consensus_history,
        running_consensus=running_consensus,
        unresolved_points=unresolved_points,
    )

    build_bazi_result_docx(
        output_path=result_output_path,
        source_prompt=source_prompt,
        final_scores=final_scores,
        final_report=final_report,
        running_consensus=running_consensus,
        unresolved_points=unresolved_points,
    )